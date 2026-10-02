#!/usr/bin/env python3
"""Blue Wolf balans konfiguratsiyasi.

Excel `blue_wolf_darajalar.xlsx` ning `Sozlamalar` varagʻi — balansning yagona manbai.
Har bir parametr Excel da nomlangan katak (game_config kaliti bilan bir xil, masalan `xp_base`),
boshqa varaqlar formulalari shu nomlar orqali hisoblanadi.
Bu skript:
  1. Sozlamalar dagi har bir parametrni `game_config` kalitiga moslaydi
     (yangi yoki oʻchirilgan qator boʻlsa — xato beradi, jimgina oʻtkazib yubormaydi);
  2. `blue_wolf_game_config.sql` ni yaratadi;
  3. balans tekshiruvlarini bajaradi (hujjatlardagi qoidalar raqamlarda bajarilayaptimi).

Ishlatish (docs/blue_wolf papkasidan):
    pip install openpyxl
    python3 tools/blue_wolf_config.py            # SQL yaratadi + tekshiradi
    python3 tools/blue_wolf_config.py --check    # faqat tekshiradi, SQL ni oʻzgartirmaydi
"""
import argparse
import math
import pathlib
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
XLSX = ROOT / "blue_wolf_darajalar.xlsx"
SQL = ROOT / "blue_wolf_game_config.sql"

# Sozlamalar varagʻidagi yorliq -> game_config kaliti
KEYS = {
    "Bazaviy XP (1→2 daraja)": "xp_base",
    "XP oʻsish koeffitsienti": "xp_growth",
    "Bazaviy kuch": "stat_power_base",
    "Kuch oʻsishi": "stat_power_growth",
    "Bazaviy tezlik": "stat_speed_base",
    "Tezlik oʻsishi": "stat_speed_growth",
    "Bazaviy chidam": "stat_hp_base",
    "Chidam oʻsishi": "stat_hp_growth",
    "Haqiqiy boʻrilar": "class_real",
    "Qadimgi boʻrilar": "class_ancient",
    "Mifologik boʻrilar": "class_myth",
    "Toʻda ochilish darajasi": "pack_unlock_level",
    "Erta bosqich koeffitsienti": "stage_early",
    "Oʻrta bosqich koeffitsienti": "stage_mid",
    "Oʻrta bosqich chegarasi": "stage_mid_level",
    "Oddiy bosqich koeffitsienti": "stage_normal",
    "Kech bosqich koeffitsienti": "stage_late",
    "Kech bosqich chegarasi": "stage_late_level",
    "Tanishtiruv tugash darajasi": "tutorial_end_level",
    "Tanishtiruv maqsadi (daqiqa)": "tutorial_target_min",
    "Bepul tezlashtirish": "tutorial_free_speedups",
    "Oʻtkazib yuborish qadami": "tutorial_skip_step",
    "Kundalik vazifa soni": "quest_daily_count",
    "Haftalik vazifa soni": "quest_weekly_count",
    "Kundalik mukofot chegarasi": "quest_daily_cap",
    "Haftalik mukofot chegarasi": "quest_weekly_cap",
    "Bosqichli mukofot chegarasi": "quest_milestone_cap",
    "Tezlashtirish mukofoti": "quest_speedup_min",
    "Xavfsiz chegara": "quest_safe_cap",
    "Yangi askar: bazaviy goʻsht": "train_meat_base",
    "Yangi askar: bazaviy suyak": "train_bone_base",
    "Yangi askar: bazaviy vaqt": "train_time_min",
    "Almashtirish yoʻqotishi": "promote_loss",
    "Almashtirish narx ulushi": "promote_cost_share",
    "Almashtirish vaqt ulushi": "promote_time_share",
    "Normal toʻlganlik": "occupancy_normal",
    "Toʻlganlik eksponenti": "occupancy_exp",
    "Minimal vaqt koeffitsienti": "occupancy_min",
    "Yakuniy koeff. chegarasi": "train_coef_max",
    "Oflayn ombor kengaytmasi": "offline_store_mult",
    "Oflayn ishlab chiqarish": "offline_prod_rate",
    "Ochlik CP jazosi": "hunger_cp_penalty",
    "Ochlik ishlab chiqarish jazosi": "hunger_prod_penalty",
    "Ochlikda ketish boshlanishi": "hunger_leave_after_d",
    "Kunlik ketish ulushi": "hunger_leave_daily",
    "Maksimal ochlik yoʻqotishi": "hunger_leave_max",
    "Uyqu qalqoni": "sleep_shield_h",
    "Qaytish grace vaqti": "comeback_shield_min",
    "Kunlik tezlashtirish chegarasi": "speedup_daily_cap",
    "Tezlashtirish bazaviy narxi": "speedup_base_price",
    "Tezlashtirish narx oʻsishi": "speedup_price_growth",
    "Blok uzunligi": "speedup_block_h",
    "Oy toshi kursi": "moonstone_per_usd",
    "Mavsum yoʻli narxi": "price_season_pass",
    "Ikkinchi navbat narxi": "price_second_queue",
    "Qaytish narx koeffitsienti": "return_price_coef",
    "Qaytish maksimal qisqartirish": "return_speedup_max",
    "Taʼtil kechikishi": "vacation_delay_h",
    "Taʼtil minimal muddati": "vacation_min_h",
    "Taʼtil maksimal muddati": "vacation_max_h",
    "Bepul taʼtil": "vacation_free_h",
    "Taʼtil narxi": "vacation_price_day",
    "Taklif etiladigan raqib": "match_offers",
    "Roʻyxat yangilanishi": "match_refresh_min",
    "Bazaviy yoʻqotish": "battle_base_loss",
    "Maksimal yoʻqotish": "battle_max_loss",
    "Jang raundlari": "battle_rounds",
    "Tasodif oraligʻi": "battle_variance",
    "Hujumchi oʻlim ulushi": "att_death_share",
    "Himoyachi oʻlim ulushi": "def_death_share",
    "Yurish tezligi": "march_speed",
    "Razvedka tezligi": "scout_speed",
    "Toʻda jangi davomiyligi": "war_duration_h",
    "Klan kuch oynasi": "war_power_window",
    "Minimal ishtirok": "war_min_participation",
    "Minimal aʼzo ishtiroki": "war_min_members",
    "Takroriy urush bloki": "war_repeat_block_d",
    "Urush garovi": "war_deposit",
    "Gʻolib yoʻqotishi": "war_winner_loss",
    "Magʻlub yoʻqotishi": "war_loser_loss",
    "Urush oʻlim ulushi (gʻolib)": "war_winner_death_share",
    "Urush oʻlim ulushi (magʻlub)": "war_loser_death_share",
    "Aʼzo mukofot shifti": "war_member_cap",
    "Magʻlub ishtirok mukofoti": "war_loser_reward_share",
    "Klan XP bazasi": "clan_xp_base",
    "Klan XP oʻsishi": "clan_xp_growth",
    "Yolgʻiz bosqich toʻdasi": "pack_solo_size",
    "Bazaviy toʻda hajmi": "pack_base",
    "Toʻda oʻsish koeffitsienti": "pack_growth",
    "CP: kuch vazni": "cp_w_power",
    "CP: tezlik vazni": "cp_w_speed",
    "CP: chidam vazni": "cp_w_hp",
    "Bazaviy kunlik ehtiyoj": "need_base",
    "Ehtiyoj oʻsishi": "need_growth",
    "Bazaviy suv ehtiyoji": "water_need_base",
    "Suv oʻsishi": "water_need_growth",
    "Oy nuri ehtiyoji": "moonlight_need",
    "Zaxira kuni": "reserve_days",
    "Gʻor bazaviy sigʻimi": "store_base",
    "Gʻor sigʻim oʻsishi": "store_growth",
    "Bazaviy passiv suv": "water_passive_base",
    "Passiv suv oʻsishi": "water_passive_growth",
    "Bazaviy himoya": "cave_protect_base",
    "Himoya oʻsishi": "cave_protect_growth",
    "Gʻor narxi: tosh": "cost_food_cave_stone",
    "Gʻor narxi: shox-shabba": "cost_food_cave_wood",
    "Bazaviy ishlab chiqarish": "prod_base",
    "Ishlab chiqarish oʻsishi": "prod_growth",
    "Ustaxona buferi": "workshop_buffer_h",
    "Ustaxona narxi: tosh": "cost_workshop_stone",
    "Ustaxona narxi: shox-shabba": "cost_workshop_wood",
    "Ustaxona narxi: suyak": "cost_workshop_bone",
    "Bino narx oʻsishi": "cost_growth",
    "Bino vaqt oʻsishi": "time_growth",
    "Bazaviy qurilish vaqti": "build_time_base_min",
    "Kuchaytirish rejasi": "build_plan_days",
    "Alfa bonusi": "alpha_bonus",
    "Tier koeffitsienti": "tier_coef",
    "Tier ochilish qadami": "tier_step_level",
    "Qarshi-kuch bonusi": "counter_bonus",
    "Namuna alfa darajasi": "sample_alpha_level",
    "Taqsimot: razvedkachi": "share_scout",
    "Taqsimot: hujumchi": "share_attacker",
    "Taqsimot: himoyachi": "share_defender",
    "Taqsimot: ovchi": "share_hunter",
    "Tier qadami (bino)": "tier_step_building",
    "Bino cheklovi": "building_level_over_player",
    "Maydon narxi: tosh": "cost_role_stone",
    "Maydon narxi: suyak": "cost_role_bone",
    "Maydon narxi: teri": "cost_role_hide",
    "Mashq tezligi bonusi": "train_speed_per_level",
    "Bino narxi: Razvedka qoyasi": "cost_coef_scout_rock",
    "Bino narxi: Jang maydoni": "cost_coef_battle_ground",
    "Bino narxi: Himoya devori": "cost_coef_defense_wall",
    "Bino narxi: Ov soʻqmogʻi": "cost_coef_hunt_path",
    "Lager yigʻim koeffitsienti": "camp_rate",
    "Lager tier bonusi": "camp_tier_bonus",
    "Lager muddati": "camp_hours",
    "Lagerga yuborish ulushi": "camp_send_share",
    "Oazis soni (server)": "oasis_count",
    "Oazis bonusi": "oasis_bonus_water",
    "Oazis ushlab turish": "oasis_hold_h",
    "Bazaviy askar sigʻimi": "role_cap_base",
    "Sigʻim oʻsish koeffitsienti": "role_cap_growth",
    "Ortiqcha jazo darajasi": "role_cap_penalty_exp",
    "Maksimal jazo": "role_cap_penalty_max",
    "Vaqt: Oziq gʻori": "time_coef_food_cave",
    "Vaqt: Ustaxona qoyasi": "time_coef_workshop",
    "Vaqt: Razvedka qoyasi": "time_coef_scout_rock",
    "Vaqt: Jang maydoni": "time_coef_battle_ground",
    "Vaqt: Himoya devori": "time_coef_defense_wall",
    "Vaqt: Ov soʻqmogʻi": "time_coef_hunt_path",
    "Vaqt: Shifo gʻori": "time_coef_hospital",
    "Vaqt: Bozor": "time_coef_market",
    "Shifo bazaviy sigʻimi": "hospital_cap_base",
    "Shifo sigʻim oʻsishi": "hospital_cap_growth",
    "Bazaviy davolash vaqti": "heal_time_min",
    "Davolash tezligi oʻsishi": "heal_speed_growth",
    "Shifo narxi: shox-shabba": "cost_hospital_wood",
    "Shifo narxi: teri": "cost_hospital_hide",
    "Shifo narxi: tosh": "cost_hospital_stone",
    "Bazaviy almashinuv kursi": "market_rate_base",
    "Kurs yaxshilanishi": "market_rate_step",
    "Bazaviy kunlik limit": "market_limit_base",
    "Limit oʻsishi": "market_limit_growth",
    "Bazaviy soliq": "market_tax_base",
    "Soliq kamayishi": "market_tax_step",
    "Bozor narxi: tosh": "cost_market_stone",
    "Bozor narxi: shox-shabba": "cost_market_wood",
    "Faol oʻyinchi soni": "fund_dau_estimate",
    "Kunlik ARPU": "fund_arpu_estimate",
    "Jamgʻarma ulushi": "fund_share",
    "Mavsum davomiyligi": "season_days",
    "Maksimal shaxsiy ulush": "season_player_cap",
    "Minimal haftalik janglar": "season_min_battles",
    "Yakuniy kunlar bonusi": "season_final_days_mult",
    "Tanga kursi": "coins_per_usd",
    "Yangi akkaunt kutishi": "season_new_account_d",
    "Juftlik hissa chegarasi": "fair_pair_season_limit",
    "Yopiq guruh chegarasi": "fair_closed_group_share",
    "Yopiq guruh hajmi": "fair_closed_group_size",
    "Minimal himoya CP": "fair_min_defense_cp",
    "Minimal raund": "fair_min_rounds",
    "Minimal hujumchi yoʻqotishi": "fair_min_att_loss",
    "Bir tomonlama nisbat": "fair_one_sided_ratio",
    "Toʻda almashtirish bloki": "fair_clan_switch_d",
    "Toʻlov kutish vaqti": "fair_payout_wait_h",
    "Qoʻlda tekshiruv": "fair_manual_top",
    "Bazaviy yuk": "loot_carry_base",
    "Tier yuk bonusi": "loot_carry_tier",
    "Oʻlja chegarasi": "loot_cap_ratio",
    "Reyd koeffitsienti": "raid_coef",
    "Hujum oynasi: past": "attack_window_down",
    "Hujum oynasi: yuqori": "attack_window_up",
    "Qalqon darajasi": "shield_newbie_level",
    "Kengayish kutish vaqti": "window_expand_min",
    "Maksimal oyna": "window_max",
    "Minimal havza": "window_min_pool",
    "Qisman chegara": "scout_partial",
    "Toʻliq chegara": "scout_full",
    "Aniq chegara": "scout_exact",
    "Razvedka kutish vaqti": "scout_cooldown_min",
    "Maʼlumot muddati": "scout_report_min",
    "Muvaffaqiyatsiz jarohat": "scout_fail_injury",
    "Oʻlja aniqligi": "scout_loot_accuracy",
    "Qasos bonusi": "revenge_bonus",
    "Ov XP": "xp_hunt_coef",
    "Qurilish XP": "xp_build_coef",
    "PvP XP": "xp_pvp_coef",
    "PvP XP: gʻalaba": "xp_pvp_win",
    "PvP XP: durang": "xp_pvp_draw",
    "PvP XP: magʻlubiyat": "xp_pvp_loss",
    "Vazifa XP ulushi": "xp_quest_share",
    "Catch-up maksimal qoʻshimcha": "catchup_max_mult",
    "Catch-up koeffitsienti": "catchup_gap_coef",
    "Catch-up chegarasi": "catchup_gap_threshold",
    "Catch-up pasayish boshi": "catchup_taper_start",
    "Catch-up pasayish oxiri": "catchup_taper_end",
    "Ovchi ochilish darajasi": "hunter_unlock_level",
    "Yolgʻiz ov maksimal darajasi": "solo_hunt_max_level",
    "Yolgʻiz ov kutish vaqti": "solo_hunt_cooldown_min",
    "Ov davomiyligi": "hunt_duration_min",
    "Ovchi unumi koeffitsienti": "hunter_yield_mult",
    "Ovchi tier bonusi": "hunter_tier_bonus",
    "Oʻlja toʻdasi boʻluvchisi": "prey_kg_per_wolf",
    "Kichik oʻlja jazosi": "hunt_small_party_penalty",
    "Ov guruhi bonusi": "hunt_party_bonus",
    "Ov guruhi maksimal hajmi": "hunt_party_max",
    "Ovdan shifobaxsh oʻt": "hunt_herb_share",
    "Oy nuri ehtiyoj darajasi": "moonlight_need_level",
    "Oy nuri passiv ishlab chiqarishi": "moonlight_passive_base",
    "Oy mehrobi bonusi": "moon_altar_bonus",
    "Oy nuri yetishmasligi CP jazosi": "moonlight_cp_penalty",
    "Bozor: oy nuri kurs koeffitsienti": "market_moonlight_mult",
    "Suv yetishmasligi tezlik jazosi": "thirst_speed_penalty",
    "Davolash narxi: oʻt": "heal_herb_per_tier",
    "Shifo gʻorisiz tuzalish vaqti": "heal_no_hospital_min",
    "Toʻliq gʻalaba chegarasi": "battle_full_win_r",
    "Gʻalaba chegarasi": "battle_win_r",
    "Magʻlubiyat chegarasi": "battle_loss_r",
    "Tuzoq ehtimoli": "trap_chance",
    "Tuzoq zarari": "trap_damage",
    "Chekinish chegarasi": "retreat_loss",
    "Himoya devori bonusi": "wall_bonus_per_level",
    "Bir raqibga kunlik hujum": "pair_limit",
    "Qisman gʻalaba oʻlja koeffitsienti": "loot_partial_win",
    "Takror hujum koeff. (2-hujum)": "loot_repeat_2",
    "Takror hujum koeff. (3-hujum)": "loot_repeat_3",
    "Daraja oʻlja koeff. −1": "loot_level_m1",
    "Daraja oʻlja koeff. +1": "loot_level_p1",
    "Daraja oʻlja koeff. −2": "loot_level_m2",
    "Daraja oʻlja koeff. +2": "loot_level_p2",
    "Daraja oʻlja koeff. −3": "loot_level_m3",
    "Daraja oʻlja koeff. +3": "loot_level_p3",
    "Oflayn himoyachi oʻlja koeff.": "loot_offline_mult",
    "Oflayn oʻlja chegarasi": "loot_offline_after_h",
    "Botlar ochilish darajasi": "bot_unlock_level",
    "Bot kuchi: kuchsiz": "bot_power_weak",
    "Bot kuchi: teng": "bot_power_even",
    "Bot kuchi: kuchli": "bot_power_strong",
    "Bot tasodif oraligʻi": "bot_variance",
    "Bekor qilishda qaytarish": "queue_cancel_refund",
    "Ikkinchi navbat bepul darajasi": "second_queue_free_level",
    "Oflayn ombor xaridi": "offline_store_plus",
    "Oflayn chegarasi": "offline_after_min",
    "In mashq tezlanishi": "den_speed_coef",
    "Lager ulushi tavani": "camp_yield_share_max",
    "Reyd qalqoni: yoʻqotish chegarasi": "shield_after_raid_loss",
    "Reyd qalqoni": "shield_after_raid_h",
    "Ketma-ket magʻlubiyat qalqoni": "shield_two_losses_h",
    "Tanishtiruv tezlashtirish uzunligi": "tutorial_speedup_min",
    "Urushni rad etish jarimasi": "war_decline_penalty",
    "Tungi reyd tezlik bonusi": "night_raid_speed_bonus",
}


def read_settings(wb):
    ws = wb["Sozlamalar"]
    params = {}
    for r in range(1, ws.max_row + 1):
        label, value, unit, note = (ws.cell(r, c).value for c in range(1, 5))
        if not isinstance(label, str) or not isinstance(value, (int, float)):
            continue
        if label not in KEYS:
            sys.exit(f"Sozlamalar!A{r}: '{label}' uchun game_config kaliti yoʻq — KEYS ga qoʻshing")
        key = KEYS[label]
        if key in params:
            sys.exit(f"Takroriy kalit: {key}")
        params[key] = (float(value), unit, note or label)
    missing = sorted(set(KEYS.values()) - set(params))
    if missing:
        sys.exit("Sozlamalar varagʻida topilmadi: " + ", ".join(missing))
    return params


def write_sql(params):
    def q(s):
        return "NULL" if s is None else "'" + str(s).replace("'", "''") + "'"

    rows = [f"  ({q(k)}, {v:g}, {q(u)}, {q(n)})" for k, (v, u, n) in params.items()]
    SQL.write_text(
        "-- AVTOMATIK YARATILGAN — qoʻlda tahrirlamang.\n"
        "-- Manba: blue_wolf_darajalar.xlsx, “Sozlamalar” varagʻi.\n"
        "-- Qayta yaratish: python3 tools/blue_wolf_config.py\n\n"
        "USE bluewolf;\n\n"
        "INSERT INTO game_config (config_key, config_value, unit, note) VALUES\n"
        + ",\n".join(rows)
        + "\nAS new ON DUPLICATE KEY UPDATE\n"
        "  config_value = new.config_value, unit = new.unit, note = new.note;\n",
        encoding="utf-8",
    )


class Balance:
    """Hujjatlardagi formulalar — tekshiruvlar uchun."""

    def __init__(self, p):
        self.p = {k: v[0] for k, v in p.items()}

    def __getattr__(self, k):
        return self.p[k]

    def stage(self, L):
        if L <= 4:
            return self.stage_early
        if L <= self.stage_mid_level:
            return self.stage_mid
        if L <= self.stage_late_level:
            return self.stage_normal
        return self.stage_late

    def army(self, L):
        if L < self.pack_unlock_level:
            return self.pack_solo_size
        return round(self.pack_base * self.pack_growth ** (L - self.pack_unlock_level))

    def need(self, L):
        return self.need_base + self.need_growth * (L - 1)

    def hunters(self, L):
        return max(1, math.floor(self.army(L) * self.share_hunter + 0.5))

    def max_tier(self, L, building):
        return max(1, min(6, L // int(self.tier_step_level), building // int(self.tier_step_building)))

    def cave_cap(self, L):
        return self.store_base * self.store_growth ** (L - 1)

    def cave_protect(self, L):
        return self.cave_protect_base + self.cave_protect_growth * (L - 1)

    def hunt_day(self, L, hunts=4):
        per_hunter = self.hunter_yield_mult * self.need(L)
        return hunts * self.hunters(L) * per_hunter * self.hunt_duration_min / 60


def run_checks(b, wb):
    errors = []

    def check(ok, msg):
        if not ok:
            errors.append(msg)

    # 1. Tier: har darajada kamida 1-tier, yangi bino (1-daraja) bilan ham
    for L in range(1, 26):
        check(b.max_tier(L, 1) >= 1, f"L{L}: yangi rol binosi bilan askar tayyorlab boʻlmaydi")

    # 2. Askar narxi (goʻsht) tier ochilgan darajadagi Oziq gʻori sigʻimiga sigʻadi
    for t in range(1, 7):
        L = max(1, int(b.tier_step_level) * (t - 1))
        cost = b.train_meat_base * b.tier_coef ** (2 * (t - 1))
        check(cost <= b.cave_cap(L), f"{t}-tier askar ({cost:.0f} kg) L{L} gʻorga sigʻmaydi ({b.cave_cap(L):.0f})")

    # 3. Ov: kuniga 4 ov sarfning kamida 2× ini beradi; bitta ov gʻorga sigʻadi
    for L in range(4, 26):
        cons = b.need(L) * b.army(L)
        day = b.hunt_day(L)
        check(day >= 2 * cons, f"L{L}: kunlik ov {day:.0f} < 2 × sarf {cons:.0f}")
        check(day / 4 <= b.cave_cap(L), f"L{L}: bitta ov {day/4:.0f} gʻor sigʻimidan {b.cave_cap(L):.0f} katta")

    # 4. Suv: gʻor oʻyinchi darajasida boʻlsa ehtiyojni qoplaydi
    for L in range(4, 26):
        need = (b.water_need_base + b.water_need_growth * (L - 1)) * b.army(L)
        passive = b.water_passive_base * b.water_passive_growth ** (L - 1) * 24
        check(need <= passive <= 4 * need, f"L{L}: suv {passive:.0f}/kun, ehtiyoj {need:.0f} — 1×–4× oraligʻidan tashqarida")

    # 5. Oy nuri: passiv yetmaydi, Oy mehrobi bilan yetadi
    for L in range(int(b.moonlight_need_level), 26):
        need = b.moonlight_need * b.army(L)
        passive = b.moonlight_passive_base * 24 * b.army(L)
        check(passive < need <= passive * (1 + b.moon_altar_bonus) + 1e-9,
              f"L{L}: oy nuri passiv {passive:.1f}, ehtiyoj {need:.1f}, mehrob bilan {passive*(1+b.moon_altar_bonus):.1f}")

    # 6. Oʻlja: bitta reyd kunlik ehtiyojning 50% idan oshmaydi
    for L in range(5, 26):
        stock = b.need(L) * b.army(L) * b.reserve_days
        cave_L = next(l for l in range(1, 26) if b.cave_cap(l) >= stock)
        loot = stock * (1 - b.cave_protect(cave_L)) * b.raid_coef
        cons = b.need(L) * b.army(L)
        check(loot <= b.loot_cap_ratio * cons + 0.5, f"L{L}: oʻlja {loot:.0f} > {b.loot_cap_ratio:.0%} × {cons:.0f}")

    # 7. Tanishtiruv XP: daraja ustuni qadam XP si bilan mos
    ws = wb["Tanishtiruv"]
    hdr = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Qadam")
    xp_col = next(c for c in range(1, ws.max_column + 1) if ws.cell(hdr, c).value == "XP")
    total_xp = [0, 0]
    for L in range(2, 26):
        # Darajalar varagʻi kabi: har daraja narxi 10 ga yaxlitlanadi
        total_xp.append(total_xp[-1] + round(b.xp_base * b.xp_growth ** (L - 2) * b.stage(L), -1))
    xp = 0
    for r in range(hdr + 1, hdr + 21):
        step_level = ws.cell(r, 2).value
        reached = max(L for L in range(1, 26) if xp >= total_xp[L])
        check(reached == step_level, f"Tanishtiruv {ws.cell(r, 1).value}-qadam: daraja {step_level}, XP boʻyicha {reached}")
        xp += ws.cell(r, xp_col).value
    check(xp >= total_xp[int(b.tutorial_end_level)],
          f"Tanishtiruv XP si {xp} — {int(b.tutorial_end_level)}-daraja uchun yetmaydi")

    # 8. Lager ulushi tavan bilan
    check(0 < b.camp_yield_share_max <= 0.6, "Lager tavani 0–60% oraligʻida boʻlishi kerak")

    # 9. Toʻlganlik 100% da yakuniy chegaradan oshmaydi
    occ_full = (1 / b.occupancy_normal) ** b.occupancy_exp
    check(occ_full <= b.train_coef_max, f"100% toʻlganlik koeffitsienti {occ_full:.2f} > {b.train_coef_max}")

    # 10. Jang chegaralari tartibi
    check(b.battle_loss_r < 1 < b.battle_win_r < b.battle_full_win_r, "Jang natija chegaralari tartibi notoʻgʻri")

    # 11. Tezlashtirish: kunlik maksimal narx
    hours = int(24 * b.speedup_daily_cap)
    total = sum(round(b.speedup_base_price * b.speedup_price_growth ** i) for i in range(hours))
    check(total <= 300, f"Kunlik maksimal tezlashtirish narxi {total} oy toshi — juda qimmat")

    return errors


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="faqat tekshirish")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    params = read_settings(wb)
    errors = run_checks(Balance(params), wb)
    if not args.check:
        write_sql(params)
        print(f"{SQL.name}: {len(params)} parametr yozildi")
    if errors:
        print(f"{len(errors)} ta balans xatosi:")
        for e in errors:
            print("  ✗", e)
        sys.exit(1)
    print("Balans tekshiruvlari: hammasi oʻtdi ✓")


if __name__ == "__main__":
    main()
